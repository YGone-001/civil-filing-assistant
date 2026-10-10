# File: tests/helpers.py
# Purpose: Shared, fictional-data helpers for the test suite.
# Encoding: UTF-8
"""Helpers for building wizards and inspecting generated documents.

All sample data here is deliberately fictional (names like "原告甲（测试）",
"示例证件号-…") and must never be replaced with real litigant information.
"""

from pathlib import Path

CASE_INDEX = {"loan": 0, "contract": 1, "property": 2, "labor": 3, "divorce": 4}
CASE_TYPES = tuple(CASE_INDEX)

CASE_LABELS = {
    "loan": "民间借贷纠纷",
    "contract": "买卖合同纠纷",
    "property": "物业服务合同纠纷",
    "labor": "劳动报酬/工资追索",
    "divorce": "离婚纠纷",
}


def make_wizard():
    """Return ``(wizard, presenter)`` wired together like the real app."""
    from views.main_window import LawsuitWizard
    from presenters.main_presenter import MainPresenter

    wizard = LawsuitWizard()
    presenter = MainPresenter(wizard)
    return wizard, presenter


def select_case(wizard, case_type):
    wizard.case_selection_page.case_combo.setCurrentIndex(CASE_INDEX[case_type])
    assert wizard.field("case_type") == case_type
    return case_type


def _set_party(widget, spec):
    widget["name"].setText(spec.get("name", ""))
    widget["id"].setText(spec.get("id_number", ""))
    widget["addr"].setText(spec.get("address", ""))
    widget["phone"].setText(spec.get("phone", ""))
    widget["is_company"].setChecked(spec.get("is_company", False))
    widget["credit_code"].setText(spec.get("credit_code", ""))
    widget["legal_rep"].setText(spec.get("legal_rep", ""))


def _set_party_list(page, kind, specs):
    widgets = page.plaintiff_widgets if kind == 0 else page.defendant_widgets
    for i, spec in enumerate(specs):
        if i >= len(widgets):
            page.add_party_ui(kind)
        _set_party(widgets[i], spec)


DEFAULT_PLAINTIFFS = [
    {
        "name": "原告甲（测试）",
        "id_number": "示例证件号-原告甲",
        "address": "北京市示例区示例路1号",
        "phone": "13800000000",
    }
]
DEFAULT_DEFENDANTS = [
    {
        "name": "被告乙（测试）",
        "id_number": "示例证件号-被告乙",
        "address": "上海市示例区示例街2号",
        "phone": "13900000000",
    }
]


def set_parties(wizard, plaintiffs=None, defendants=None):
    """(Re)populate the party page with the given fictional parties."""
    page = wizard.party_page
    page.initializePage()
    _set_party_list(page, 0, plaintiffs if plaintiffs is not None else DEFAULT_PLAINTIFFS)
    _set_party_list(page, 1, defendants if defendants is not None else DEFAULT_DEFENDANTS)


def fill_case_facts(wizard, case_type, confirm_flags=True):
    """Fill the case-specific "facts and reasons" page.

    ``confirm_flags=False`` leaves the tri-state factual controls (delivery,
    receipt, IOU) untouched so a test can exercise the *unknown* state.
    """
    if case_type == "loan":
        p = wizard.loan_claim_page
        p.loan_date.setText("2023年1月1日")
        p.loan_reason.setText("资金周转")
        p.principal.setText("50000")
        p.payment_method.setText("银行转账")
        p.rate.setText("14.6")
        p.start_date.setText("2023-05-01")
        if confirm_flags:
            p.iou_yes.setChecked(True)
        p.demand_date.setText("2023年10月")
        p.demand_method.setText("微信及电话")
        p.court_name.setText("上海市徐汇区人民法院")
    elif case_type == "contract":
        p = wizard.contract_claim_page
        p.contract_date.setText("2023年1月1日")
        p.contract_name.setText("产品购销合同")
        p.product_name.setText("10台测试电脑")
        p.total_amount.setText("50000")
        p.delivery_date.setText("2023年2月1日")
        if confirm_flags:
            p.del_yes.setChecked(True)
            p.sign_yes.setChecked(True)
        p.unpaid_amount.setText("30000")
        p.penalty_amount.setText("5000")
        p.overdue_start.setText("2023-03-01")
        p.overdue_end.setText("")
        p.lpr_value.setText("3.45")
        p.lpr_multiple.setText("4")
        p.court_name.setText("上海市徐汇区人民法院")
    elif case_type == "property":
        p = wizard.property_claim_page
        p.property_addr.setText("示例小区3栋1201室")
        p.house_area.setText("115.5")
        p.fee_rate.setText("2.8")
        p.period_start.setText("2024-01-01")
        p.period_end.setText("2025-12-31")
        p.late_fee_logic.setText("按合同约定计算")
        p.demand_record.setPlainText("2025年6月多次张贴催缴通知并发送短信。")
        p.total_input.setText("")
        p.court_name.setText("上海市徐汇区人民法院")
        p.update_calc()
    elif case_type == "labor":
        p = wizard.labor_claim_page
        p.emp_join_date.setText("2023-01-15")
        p.job_title.setText("测试工程师")
        p.monthly_salary.setText("8000")
        p.unpaid_months.setText("2023年12月至2024年2月")
        p.emp_term_date.setText("2024-03-20")
        p.overtime_hours.setText("40")
        p.overtime_pay.setText("3500")
        p.no_contract_no.setChecked(True)
        p.social_sec_info.setText("未按实际工资基数缴纳")
        p.court_name.setText("北京市朝阳区人民法院")
    elif case_type == "divorce":
        p = wizard.divorce_claim_page
        p.marriage_date.setText("2018-05-20")
        p.child_name.setText("测试子女")
        p.child_birthday.setText("2020-01-01")
        p.custody_preference.setText("由原告抚养")
        p.support_monthly.setText("3000")
        p.divorce_reason.setText("感情不和")
        p.separation_start_date.setText("2022-01-01")
        p.asset_description.setText("房产归原告，车辆归被告")
        p.court_name.setText("北京市朝阳区人民法院")
    return case_type


def prepare_case(wizard, case_type, plaintiffs=None, defendants=None, with_evidence=True,
                 confirm_flags=True):
    """Select a case, fill parties and facts, and load default evidence."""
    select_case(wizard, case_type)
    set_parties(wizard, plaintiffs=plaintiffs, defendants=defendants)
    fill_case_facts(wizard, case_type, confirm_flags=confirm_flags)
    if with_evidence:
        wizard.evidence_page.initializePage()
    return wizard


def fill_minimal_case(wizard, case_type, plaintiffs=None, defendants=None):
    """Select a case and fill only the required fields with fictional data.

    Optional factual fields are deliberately left empty so a test can prove that
    their absence does not produce an invented factual assertion.
    """
    select_case(wizard, case_type)
    set_parties(wizard, plaintiffs=plaintiffs, defendants=defendants)

    if case_type == "loan":
        p = wizard.loan_claim_page
        p.loan_date.setText("2023年1月1日")
        p.loan_reason.setText("资金周转")
        p.principal.setText("10000")
        p.payment_method.setText("银行转账")
        p.demand_date.setText("2023年10月")
        p.demand_method.setText("微信")
        p.court_name.setText("示例人民法院")
    elif case_type == "contract":
        p = wizard.contract_claim_page
        p.contract_date.setText("2023年1月1日")
        p.contract_name.setText("示例合同")
        p.product_name.setText("示例货物")
        p.total_amount.setText("10000")
        p.delivery_date.setText("2023年2月1日")
        p.unpaid_amount.setText("8000")
        p.court_name.setText("示例人民法院")
    elif case_type == "property":
        p = wizard.property_claim_page
        p.property_addr.setText("示例小区1栋101室")
        p.house_area.setText("100")
        p.fee_rate.setText("2.0")
        p.period_start.setText("2024-01-01")
        p.period_end.setText("2024-12-31")
        p.court_name.setText("示例人民法院")
    elif case_type == "labor":
        p = wizard.labor_claim_page
        p.emp_join_date.setText("2023-01-15")
        p.job_title.setText("示例岗位")
        p.monthly_salary.setText("6000")
        p.unpaid_months.setText("2023年12月")
        p.emp_term_date.setText("2024-03-20")
        p.court_name.setText("示例人民法院")
    elif case_type == "divorce":
        p = wizard.divorce_claim_page
        p.marriage_date.setText("2018-05-20")
        p.child_name.setText("示例子女")
        p.child_birthday.setText("2020-01-01")
        p.support_monthly.setText("2000")
        p.court_name.setText("示例人民法院")
    return wizard


def navigate_to_export(wizard):
    """Navigate a prepared wizard to the export page and return its Finish button.

    The wizard is restarted and advanced with real ``next()`` navigation so the
    test exercises the same page transitions the user triggers by clicking
    "Next".
    """
    from PySide6.QtWidgets import QWizard

    wizard.restart()
    wizard.next()  # welcome -> case selection
    wizard.next()  # case selection -> party
    wizard.next()  # party -> claim
    wizard.next()  # claim -> evidence
    wizard.next()  # evidence -> export
    assert wizard.currentId() == wizard.PAGE_EXPORT
    return wizard.button(QWizard.FinishButton)


def docx_all_text(path):
    """Return every paragraph and table cell of a .docx as one string."""
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def verify_docx_package(path):
    """Raise ``AssertionError`` unless *path* is a valid .docx package."""
    import zipfile

    path = Path(path)
    assert path.exists(), f"missing document: {path}"
    assert zipfile.is_zipfile(path), f"not a zip/docx package: {path}"
    with zipfile.ZipFile(path) as archive:
        assert "word/document.xml" in archive.namelist(), f"missing document.xml: {path}"
        assert archive.testzip() is None, f"corrupt docx package: {path}"
    return True
