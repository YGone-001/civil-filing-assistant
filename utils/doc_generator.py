# File: utils/doc_generator.py
# Purpose: Generate the docx file for lawsuit
# Encoding: UTF-8

import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
from docx.oxml.ns import qn

class DocumentGenerator:
    def __init__(self):
        self.doc = Document()
        self._set_global_style()

    def reset(self):
        self.doc = Document()
        self._set_global_style()

    def _set_global_style(self):
        for section in self.doc.sections:
            section.top_margin = Cm(2.54)
            section.bottom_margin = Cm(2.54)
            section.left_margin = Cm(3.17)
            section.right_margin = Cm(3.17)

    def _add_text(self, text, font_name, font_size, bold, align, first_line_indent):
        p = self.doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(28)
        if first_line_indent > 0:
            p.paragraph_format.first_line_indent = Pt(first_line_indent)

        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(font_size)
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        return p

    def add_title(self, text):
        self._add_text(text, "SimSun", 22, True, WD_PARAGRAPH_ALIGNMENT.CENTER, 0)

    def add_heading(self, text):
        self._add_text(text, "SimHei", 16, True, WD_PARAGRAPH_ALIGNMENT.LEFT, 0)

    def add_body(self, text, align=WD_PARAGRAPH_ALIGNMENT.JUSTIFY, first_line_indent=32):
        self._add_text(text, "FangSong_GB2312", 16, False, align, first_line_indent)

    def add_party_block(self, title, parties):
        if len(parties) == 1:
            prefixes = [f"{title}："]
        else:
            prefixes = [f"{title}{i}：" for i in range(1, len(parties) + 1)]

        for i, party in enumerate(parties):
            p = self.doc.add_paragraph()
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(28)

            run1 = p.add_run(prefixes[i])
            run1.font.name = "SimHei"
            run1._element.rPr.rFonts.set(qn('w:eastAsia'), "SimHei")
            run1.font.size = Pt(16)
            run1.bold = False

            if getattr(party, 'is_company', False):
                company_info = f"{party.name}，统一社会信用代码：{party.credit_code}，"
                if party.legal_representative:
                    company_info += f"法定代表人：{party.legal_representative}，"
                company_info += f"住所地：{party.address}，联系电话：{party.phone}。"
                run2 = p.add_run(company_info)
            else:
                run2 = p.add_run(f"{party.name}，身份证号：{party.id_number}，住址：{party.address}，联系电话：{party.phone}。")

            run2.font.name = "FangSong_GB2312"
            run2._element.rPr.rFonts.set(qn('w:eastAsia'), "FangSong_GB2312")
            run2.font.size = Pt(16)

    def add_claims(self, claims):
        self.add_heading("诉讼请求：")
        for idx, claim in enumerate(claims, 1):
            self.add_body(f"{idx}. {claim}", align=WD_PARAGRAPH_ALIGNMENT.JUSTIFY, first_line_indent=32)

    def add_facts(self, facts_text):
        self.add_heading("事实与理由：")
        self.add_body(facts_text, align=WD_PARAGRAPH_ALIGNMENT.JUSTIFY, first_line_indent=32)

    def add_footer(self, court_name):
        self._add_text("此致", "FangSong_GB2312", 16, False, WD_PARAGRAPH_ALIGNMENT.LEFT, 0)
        self._add_text(str(court_name or ""), "FangSong_GB2312", 16, False, WD_PARAGRAPH_ALIGNMENT.LEFT, 0)
        self.doc.add_paragraph()
        self._add_text("具状人：", "FangSong_GB2312", 16, False, WD_PARAGRAPH_ALIGNMENT.RIGHT, 0)
        self._add_text("年   月   日", "FangSong_GB2312", 16, False, WD_PARAGRAPH_ALIGNMENT.RIGHT, 0)

    def _set_cell_text(self, cell, text, font_name, font_size, bold):
        cell.text = ""
        p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(28)
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(font_size)
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

    def save(self, output_path):
        self.doc.save(output_path)

    def export_complaint(self, model, file_path):
        self.reset()
        self.add_title("民 事 起 诉 状")
        self.doc.add_paragraph()
        self.add_party_block("原告", model.party_manager.plaintiffs)
        self.add_party_block("被告", model.party_manager.defendants)
        self.doc.add_paragraph()
        self.add_claims(model.render_claims())
        self.doc.add_paragraph()
        self.add_facts(model.render_facts())
        self.doc.add_paragraph()
        self.doc.add_paragraph()
        self.add_footer(model.court_name)
        DocGenerator._add_pagination(self.doc)
        self.save(file_path)

    def export_evidence_list(self, model, file_path):
        self.reset()
        self.add_title("证 据 清 单")
        self.doc.add_paragraph()

        table = self.doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"

        hdr = table.rows[0].cells
        self._set_cell_text(hdr[0], "序号", "SimHei", 16, True)
        self._set_cell_text(hdr[1], "证据名称", "SimHei", 16, True)
        self._set_cell_text(hdr[2], "证明对象/目的", "SimHei", 16, True)
        self._set_cell_text(hdr[3], "备注", "SimHei", 16, True)

        for idx, ev in enumerate(getattr(model, "evidences", []) or [], 1):
            row = table.add_row().cells
            self._set_cell_text(row[0], str(idx), "FangSong_GB2312", 16, False)
            self._set_cell_text(row[1], str(ev.name or ""), "FangSong_GB2312", 16, False)
            self._set_cell_text(row[2], str(ev.target or ""), "FangSong_GB2312", 16, False)

            remark = "原件核对"
            if getattr(model, 'case_type_name', '') == "物业服务合同纠纷":
                remark = "原件核对；计算公式已在起诉状中载明"
            elif getattr(model, 'case_type_name', '') == "劳动争议":
                remark = "原件核对；欠薪计算公式已在起诉状中载明"
            self._set_cell_text(row[3], remark, "FangSong_GB2312", 16, False)

        self.save(file_path)

    def export_address_form(self, model, file_path):
        self.reset()
        self.add_title("送 达 地 址 确 认 书")
        self.doc.add_paragraph()

        self.add_body(
            "为便于人民法院依法送达诉讼文书，当事人现确认以下送达地址信息真实、准确。"
            "下列表格逐一列示向导中已填写的当事人；如有多名当事人，请分别确认各自的送达地址，"
            "不能以其中一名当事人的地址代表其他当事人。",
            align=WD_PARAGRAPH_ALIGNMENT.JUSTIFY, first_line_indent=32,
        )
        self.doc.add_paragraph()

        entries = []
        for role, members in (
            ("原告", model.party_manager.plaintiffs),
            ("被告", model.party_manager.defendants),
        ):
            for idx, party in enumerate(members, 1):
                label = role if len(members) == 1 else f"{role}{idx}"
                entries.append((label, party))

        labels = [
            "当事人",
            "姓名/名称",
            "证件号码/统一社会信用代码",
            "联系电话",
            "送达地址",
            "备注",
        ]

        if not entries:
            entries = [("（未填写当事人）", None)]

        for label, party in entries:
            if party is None:
                name = id_number = phone = addr = ""
                role_text = label
            else:
                name = getattr(party, "name", "") or ""
                if getattr(party, "is_company", False):
                    id_number = getattr(party, "credit_code", "") or ""
                else:
                    id_number = getattr(party, "id_number", "") or ""
                phone = getattr(party, "phone", "") or ""
                addr = getattr(party, "address", "") or ""
                role_text = label

            table = self.doc.add_table(rows=6, cols=2)
            table.style = "Table Grid"

            values = [
                role_text,
                name,
                id_number,
                phone,
                addr,
                "如地址、电话变更，应及时书面告知人民法院。",
            ]

            for i in range(6):
                self._set_cell_text(table.rows[i].cells[0], labels[i], "SimHei", 16, True)
                self._set_cell_text(table.rows[i].cells[1], values[i], "FangSong_GB2312", 16, False)
            self.doc.add_paragraph()

        self._add_text("确认人：", "FangSong_GB2312", 16, False, WD_PARAGRAPH_ALIGNMENT.RIGHT, 0)
        self._add_text("年   月   日", "FangSong_GB2312", 16, False, WD_PARAGRAPH_ALIGNMENT.RIGHT, 0)

        DocGenerator._add_pagination(self.doc)
        self.save(file_path)

class DocGenerator:
    @staticmethod
    def _set_paragraph_format(paragraph, font_name='FangSong_GB2312', font_size=16, bold=False, align=WD_PARAGRAPH_ALIGNMENT.LEFT, first_line_indent=0):
        paragraph.alignment = align
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(28)
        if first_line_indent > 0:
            paragraph.paragraph_format.first_line_indent = Pt(first_line_indent)

        for run in paragraph.runs:
            run.font.name = font_name
            run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
            run.font.size = Pt(font_size)
            run.bold = bold

    @staticmethod
    def _add_party_block(doc, title, parties, body_font):
        if len(parties) == 1:
            prefixes = [f"{title}："]
        else:
            prefixes = [f"{title}{i}：" for i in range(1, len(parties) + 1)]

        for i, party in enumerate(parties):
            para = doc.add_paragraph()
            run = para.add_run(prefixes[i])
            run.font.name = 'SimHei'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')
            run.font.size = Pt(16) # 三号黑体

            content_run = para.add_run(f"{party.name}，身份证号：{party.id_number}，住址：{party.address}，联系电话：{party.phone}。")
            content_run.font.name = body_font
            content_run._element.rPr.rFonts.set(qn('w:eastAsia'), body_font)
            content_run.font.size = Pt(16)

            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            para.paragraph_format.line_spacing = Pt(28)

    @staticmethod
    def _add_pagination(doc):
        from docx.oxml import OxmlElement

        for section in doc.sections:
            footer = section.footer
            para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

            run = para.add_run("- ")

            fldChar1 = OxmlElement('w:fldChar')
            fldChar1.set(qn('w:fldCharType'), 'begin')

            instrText = OxmlElement('w:instrText')
            instrText.set(qn('xml:space'), 'preserve')
            instrText.text = " PAGE "

            fldChar2 = OxmlElement('w:fldChar')
            fldChar2.set(qn('w:fldCharType'), 'separate')

            fldChar3 = OxmlElement('w:fldChar')
            fldChar3.set(qn('w:fldCharType'), 'end')

            run._r.append(fldChar1)
            run._r.append(instrText)
            run._r.append(fldChar2)
            run._r.append(fldChar3)

            para.add_run(" -")

    @staticmethod
    def generate_loan_lawsuit(model, output_path):
        gen = DocumentGenerator()
        gen.export_complaint(model, output_path)

    @staticmethod
    def generate_complaint(model, output_path):
        DocGenerator.generate_loan_lawsuit(model, output_path)

    @staticmethod
    def generate_evidence_list(model, output_path):
        gen = DocumentGenerator()
        gen.export_evidence_list(model, output_path)

    @staticmethod
    def generate_address_confirmation(model, output_path):
        gen = DocumentGenerator()
        gen.export_address_form(model, output_path)
